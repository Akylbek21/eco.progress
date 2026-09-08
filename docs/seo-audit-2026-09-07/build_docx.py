from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "report-source.md"
OUTPUT = BASE / "SEO-strategy-EcoProgress-2026-09-07.docx"

ACCENT = "244B3A"
PALE = "EEF4F0"
BORDER = "D9D9D9"
TEXT = RGBColor(0, 0, 0)
MUTED = RGBColor(82, 91, 87)


def set_font(run, name="Arial", size=None, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def suppress_paragraph_borders(ppr):
    border = ppr.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        ppr.append(border)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = border.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            border.append(node)
        node.set(qn("w:val"), "nil")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D5E48")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Arial")
    r_pr.append(r_fonts)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, size=10.5, color=TEXT):
    token_re = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True, color=color)
        elif token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if m:
                run = paragraph.add_run(m.group(1))
                set_font(run, size=size, color=RGBColor(29, 94, 72))
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=max(8, size - 1), color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=size, color=color)


def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)
        section.top_margin = Inches(0.48)
        section.bottom_margin = Inches(0.48)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.65)


def add_page_number(section):
    return


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def parse_table(lines, start):
    block = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        block.append(lines[i])
        i += 1
    rows = [[c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))] for line in block]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    if cols == 11:
        p1 = doc.add_paragraph()
        p1.paragraph_format.keep_with_next = True
        r1 = p1.add_run("Матрица кластеров — интент и решение")
        set_font(r1, size=11, bold=True)
        first = [[r[i] if i < len(r) else "" for i in (0, 1, 2, 3, 4, 6, 10)] for r in rows]
        add_table(doc, first)
        p2 = doc.add_paragraph()
        p2.paragraph_format.keep_with_next = True
        r2 = p2.add_run("Матрица кластеров — URL и метаданные")
        set_font(r2, size=11, bold=True)
        second = [[r[i] if i < len(r) else "" for i in (0, 5, 7, 8, 9)] for r in rows]
        add_table(doc, second)
        return
    wide = cols >= 8
    if wide:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, landscape=True)
        add_page_number(section)

    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    font_size = 6.7 if cols >= 9 else 8.2 if cols >= 6 else 9.0

    if cols == 11:
        weights = [0.35, 0.85, 0.9, 0.82, 0.88, 0.72, 0.52, 0.7, 1.38, 1.05, 0.43]
    elif cols == 7:
        weights = [1.5, 0.65, 0.55, 0.65, 0.65, 0.55, 1.35]
    elif cols == 6:
        weights = [1.25, 1.25, 1.0, 1.1, 1.15, 1.55]
    elif cols == 5:
        weights = [0.35, 1.1, 0.7, 1.15, 2.9]
    elif cols == 4:
        weights = [0.35, 2.35, 0.85, 1.45]
    elif cols == 3:
        weights = [1.0, 0.9, 3.6]
    else:
        weights = [1.0] * cols
    available = 10.0 if wide else 6.85
    unit = available / sum(weights)

    for ri, values in enumerate(rows):
        row = table.add_row()
        if ri == 0:
            repeat_header(row)
        for ci, cell in enumerate(row.cells):
            value = values[ci] if ci < len(values) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.width = Inches(weights[ci] * unit)
            set_cell_border(cell)
            set_cell_margins(cell, 70 if wide else 95, 70 if wide else 95, 70 if wide else 95, 70 if wide else 95)
            if ri == 0:
                shade(cell, ACCENT)
            elif ri % 2 == 0:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 0 or len(value) < 14 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=font_size, color=RGBColor(255,255,255) if ri == 0 else TEXT)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)

    if wide:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, landscape=False)
        add_page_number(section)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = TEXT
    title.paragraph_format.space_after = Pt(18)
    suppress_paragraph_borders(title._element.get_or_add_pPr())

    for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 14, 14, 7), ("Heading 3", 11.5, 10, 5)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEXT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Source URL" not in styles:
        s = styles.add_style("Source URL", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Arial"
        s.font.size = Pt(8.5)
        s.font.color.rgb = MUTED
        s.paragraph_format.left_indent = Inches(0.18)
        s.paragraph_format.space_after = Pt(4)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_section(doc.sections[0], landscape=False)
    add_page_number(doc.sections[0])
    setup_styles(doc)

    i = 0
    first_h2 = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            suppress_paragraph_borders(p._p.get_or_add_pPr())
            add_inline(p, line[2:], size=28)
        elif line.startswith("## "):
            heading_text = line[3:]
            major_break = bool(re.match(r"(?:Executive Summary|1\.|4\.|5\.|6\.|10\.|12\.|13\.|14\.|15\.)", heading_text))
            if first_h2 or major_break:
                doc.add_page_break()
                first_h2 = False
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, heading_text, size=18)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], size=14)
        elif line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:], size=11.5)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    core = doc.core_properties
    core.title = "SEO стратегия EcoProgress.kz"
    core.subject = "B2B экологические услуги в Казахстане"
    core.author = "SEO research report"
    core.keywords = "EcoProgress, SEO, Казахстан, экологические услуги"
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
